import csv
import gc
import io
import json
import os
import re
import subprocess
import tempfile
import zipfile
from collections.abc import Callable, Iterator, Sequence
from email.parser import Parser as EmailParser
from io import BytesIO
from pathlib import Path
from typing import IO, TYPE_CHECKING, Any, NamedTuple, Optional, cast
from zipfile import BadZipFile

import chardet
import openpyxl
from openpyxl.worksheet._read_only import ReadOnlyWorksheet

from onyx.configs.app_configs import (
    MAX_EMBEDDED_IMAGES_PER_FILE,
    MAX_XLSX_CELLS_PER_SHEET,
    PDF_TEXT_EXTRACTION_TIMEOUT_SECONDS,
)
from onyx.configs.constants import ONYX_METADATA_FILENAME
from onyx.configs.llm_configs import get_image_extraction_and_analysis_enabled
from onyx.file_processing.file_types import (
    PRESENTATION_MIME_TYPE,
    WORD_PROCESSING_MIME_TYPE,
    OnyxFileExtensions,
    OnyxMimeTypes,
)
from onyx.file_processing.html_utils import parse_html_page_basic
from onyx.file_processing.pdf_image_utils import iter_pdf_extracted_images
from onyx.file_processing.unstructured import (
    get_unstructured_api_key,
    unstructured_to_text,
)
from onyx.utils.logger import setup_logger
from onyx.utils.process_isolation import IsolatedProcessError, run_in_isolated_process

if TYPE_CHECKING:
    from markitdown import MarkItDown
logger = setup_logger()

TEXT_SECTION_SEPARATOR = "\n\n"

# Legacy binary Word (.doc) extraction via the antiword subprocess.
DOC_TEXT_EXTRACTION_TIMEOUT_SECONDS = 60
# Cap on expanded `number-columns-repeated` cells in ODF tables; larger repeats
# are trailing-formatting padding and get trimmed downstream anyway.
MAX_ODS_CELL_REPEAT = 64

_MARKITDOWN_CONVERTER: Optional["MarkItDown"] = None

KNOWN_OPENPYXL_BUGS = [
    "Value must be either numerical or a string containing a wildcard",
    "File contains no valid workbook part",
    "Unable to read workbook: could not read stylesheet from None",
    "Colors must be aRGB hex values",
    "Max value is",
    "There is no item named",
]


def get_markitdown_converter() -> "MarkItDown":
    global _MARKITDOWN_CONVERTER

    if _MARKITDOWN_CONVERTER is None:
        from markitdown import MarkItDown

        # Patch this function to effectively no-op because we were seeing this
        # module take an inordinate amount of time to convert charts to markdown,
        # making some powerpoint files with many or complicated charts nearly
        # unindexable.
        from markitdown.converters._pptx_converter import PptxConverter

        setattr(  # noqa: B010
            PptxConverter,
            "_convert_chart_to_markdown",
            lambda self, chart: "\n\n[chart omitted]\n\n",  # noqa: ARG005
        )
        _MARKITDOWN_CONVERTER = MarkItDown(enable_plugins=False)
    return _MARKITDOWN_CONVERTER


def get_file_ext(file_path_or_name: str | Path) -> str:
    _, extension = os.path.splitext(file_path_or_name)
    return extension.lower()


def is_text_file(file: IO[bytes]) -> bool:
    """
    checks if the first 1024 bytes only contain printable or whitespace characters
    if it does, then we say it's a plaintext file
    """
    raw_data = file.read(1024)
    file.seek(0)
    text_chars = bytearray({7, 8, 9, 10, 12, 13, 27} | set(range(0x20, 0x100)) - {0x7F})
    return all(c in text_chars for c in raw_data)


def detect_encoding(file: IO[bytes]) -> str:
    """Detect the character encoding of a binary file.

    Tries UTF-8 first — if the bytes decode cleanly, they are definitively UTF-8
    (UTF-8 is self-validating). Falls back to chardet only when UTF-8 fails, since
    chardet can misidentify valid UTF-8 text (e.g. Cyrillic) as a legacy encoding
    like windows-1251, producing mojibake. Defaults to utf-8 if chardet gives up.

    Resets the file cursor to 0 after sampling so callers can still read the full file.
    """
    raw_data = file.read(50000)
    file.seek(0)
    try:
        raw_data.decode("utf-8")
        return "utf-8"
    except UnicodeDecodeError:
        # utf-8 failed — bytes are genuinely non-UTF-8, let chardet guess
        encoding = chardet.detect(raw_data)["encoding"] or "utf-8"
        return encoding


def is_macos_resource_fork_file(file_name: str) -> bool:
    return os.path.basename(file_name).startswith("._") and file_name.startswith(
        "__MACOSX"
    )


def to_bytesio(stream: IO[bytes]) -> BytesIO:
    if isinstance(stream, BytesIO):
        return stream
    data = stream.read()  # consumes the stream!
    return BytesIO(data)


def load_files_from_zip(
    zip_file_io: IO,
    ignore_macos_resource_fork_files: bool = True,
    ignore_dirs: bool = True,
) -> Iterator[tuple[zipfile.ZipInfo, IO[Any]]]:
    """
    Iterates through files in a zip archive, yielding (ZipInfo, file handle) pairs.
    """
    with zipfile.ZipFile(zip_file_io, "r") as zip_file:
        for file_info in zip_file.infolist():
            if ignore_dirs and file_info.is_dir():
                continue

            if (
                ignore_macos_resource_fork_files
                and is_macos_resource_fork_file(file_info.filename)
            ) or file_info.filename == ONYX_METADATA_FILENAME:
                continue

            with zip_file.open(file_info.filename, "r") as subfile:
                # Try to match by exact filename first
                yield file_info, subfile


def _extract_onyx_metadata(line: str) -> dict | None:
    """
    Example: first line has:
        <!-- ONYX_METADATA={"title": "..."} -->
      or
        #ONYX_METADATA={"title":"..."}
    """
    html_comment_pattern = r"<!--\s*ONYX_METADATA=\{(.*?)\}\s*-->"
    hashtag_pattern = r"#ONYX_METADATA=\{(.*?)\}"

    html_comment_match = re.search(html_comment_pattern, line)
    hashtag_match = re.search(hashtag_pattern, line)

    if html_comment_match:
        json_str = html_comment_match.group(1)
    elif hashtag_match:
        json_str = hashtag_match.group(1)
    else:
        return None

    try:
        return json.loads("{" + json_str + "}")
    except json.JSONDecodeError:
        return None


def read_text_file(
    file: IO,
    encoding: str = "utf-8",
    errors: str = "replace",
    ignore_onyx_metadata: bool = True,
) -> tuple[str, dict]:
    """
    For plain text files. Optionally extracts Onyx metadata from the first line.
    """
    metadata = {}
    file_content_raw = ""
    for ind, line in enumerate(file):
        # decode
        try:
            line = line.decode(encoding) if isinstance(line, bytes) else line
        except UnicodeDecodeError:
            line = (
                line.decode(encoding, errors=errors)
                if isinstance(line, bytes)
                else line
            )

        # optionally parse metadata in the first line
        if ind == 0 and not ignore_onyx_metadata:
            potential_meta = _extract_onyx_metadata(line)
            if potential_meta is not None:
                metadata = potential_meta
                continue

        file_content_raw += line

    return file_content_raw, metadata


def pdf_to_text(
    file: IO[Any], pdf_pass: str | None = None, isolate_pdfium: bool = True
) -> str:
    """
    Extract text from a PDF. For embedded images, a more complex approach is needed.
    This is a minimal approach returning text only.
    """
    text, _, _ = read_pdf_file(file, pdf_pass, isolate_pdfium=isolate_pdfium)
    return text


def _extract_pdf_text_pdfium(file_bytes: bytes, password: str | None) -> str:
    """Extract all text from a PDF via pypdfium2 (PDFium/C).

    PDFium releases the GIL while parsing, so a large or complex PDF can't pin
    a worker thread or stall the indexing heartbeat during text extraction.
    """
    import pypdfium2 as pdfium

    pdf = pdfium.PdfDocument(file_bytes, password=password)
    try:
        page_texts: list[str] = []
        for page in pdf:
            # Per-page try/finally so a get_textpage() failure still closes the
            # native page handle instead of leaking it until GC.
            try:
                textpage = page.get_textpage()
                try:
                    page_texts.append(textpage.get_text_range())
                finally:
                    textpage.close()
            finally:
                page.close()
        return TEXT_SECTION_SEPARATOR.join(page_texts)
    finally:
        pdf.close()


def read_pdf_file(
    file: IO[Any],
    pdf_pass: str | None = None,
    extract_images: bool = False,
    image_callback: Callable[[bytes, str], None] | None = None,
    isolate_pdfium: bool = True,
) -> tuple[str, dict[str, Any], Sequence[tuple[bytes, str]]]:
    """
    Returns the text, basic PDF metadata, and optionally extracted images.

    ``isolate_pdfium=False`` is for a caller that is itself a child process
    under a deadline: a child of a child is orphaned when the outer one is
    killed, so PDFium runs in the caller's process instead.
    """
    from pypdf import PdfReader
    from pypdf.errors import PdfStreamError
    from pypdfium2 import PdfiumError

    metadata: dict[str, Any] = {}
    extracted_images: list[tuple[bytes, str]] = []
    try:
        # Read once: the text extractor and the metadata/image reader share these bytes.
        file_bytes = file.read()
        pdf_reader = PdfReader(io.BytesIO(file_bytes))

        decrypt_password: str | None = None
        if pdf_reader.is_encrypted:
            # Try the explicit password first, then fall back to an empty
            # string.  Owner-password-only PDFs (permission restrictions but
            # no open password) decrypt successfully with "".
            # See https://github.com/onyx-dot-app/onyx/issues/9754
            passwords = [p for p in [pdf_pass, ""] if p is not None]
            decrypt_success = False
            for pw in passwords:
                try:
                    if pdf_reader.decrypt(pw) != 0:
                        decrypt_success = True
                        decrypt_password = pw
                        break
                except Exception:
                    pass

            if not decrypt_success:
                logger.error(
                    "Encrypted PDF could not be decrypted, returning empty text."
                )
                return "", metadata, []

        # Basic PDF metadata
        if pdf_reader.metadata is not None:
            for key, value in pdf_reader.metadata.items():
                clean_key = key.lstrip("/")
                if isinstance(value, str) and value.strip():
                    metadata[clean_key] = value
                elif isinstance(value, list) and all(
                    isinstance(item, str) for item in value
                ):
                    metadata[clean_key] = ", ".join(value)

        # PDFium can hard-abort or hang on a malformed PDF (uncatchable in-process),
        # so run it isolated; a crash, timeout, or PdfiumError falls back to pypdf.
        try:
            if isolate_pdfium:
                text = run_in_isolated_process(
                    _extract_pdf_text_pdfium,
                    file_bytes,
                    decrypt_password,
                    timeout=PDF_TEXT_EXTRACTION_TIMEOUT_SECONDS,
                )
            else:
                text = _extract_pdf_text_pdfium(file_bytes, decrypt_password)
        except (PdfiumError, IsolatedProcessError) as pdfium_err:
            logger.warning(
                "PDFium text extraction failed (%s); falling back to pypdf",
                pdfium_err,
            )
            text = TEXT_SECTION_SEPARATOR.join(
                page.extract_text() for page in pdf_reader.pages
            )

        if extract_images:
            images = iter_pdf_extracted_images(pdf_reader, MAX_EMBEDDED_IMAGES_PER_FILE)
            if image_callback is None:
                extracted_images.extend(images)
            else:
                # Stream each image out immediately
                for img_bytes, image_name in images:
                    image_callback(img_bytes, image_name)

        return text, metadata, extracted_images

    except PdfStreamError as e:
        # Malformed/truncated PDF content — a per-document content issue, not
        # a platform bug. The function returns empty text and the connector
        # continues with the next doc; no need to ship a stack trace to
        # Sentry for every corrupt file we encounter.
        logger.warning("Invalid PDF file, skipping content extraction: %s", e)
    except Exception as e:
        # Unknown PDF parsing failure — elevate just the message, not a
        # traceback, for the same reason. Callers treat empty text as a
        # non-fatal skip.
        logger.warning("Failed to read PDF, skipping content extraction: %s", e)

    return "", metadata, []


def extract_docx_images(docx_bytes: IO[Any]) -> Iterator[tuple[bytes, str]]:
    """
    Given the bytes of a docx file, extract all the images.
    Returns a list of tuples (image_bytes, image_name).
    """
    try:
        with zipfile.ZipFile(docx_bytes) as z:
            for name in z.namelist():
                if name.startswith("word/media/"):
                    yield (z.read(name), name.split("/")[-1])
    except Exception:
        logger.exception("Failed to extract all docx images")


def count_docx_embedded_images(file: IO[Any], cap: int) -> int:
    """Return the number of embedded images in a docx, short-circuiting at cap+1.

    Mirrors count_pdf_embedded_images so upload validation can apply the same
    per-file/per-batch caps. Returns a value > cap once the count exceeds the
    cap so callers do not iterate every media entry just to report a number.
    Always restores the file pointer to its original position before returning.
    """
    try:
        start_pos = file.tell()
    except Exception:
        start_pos = None
    try:
        if start_pos is not None:
            file.seek(0)
        count = 0
        with zipfile.ZipFile(file) as z:
            for name in z.namelist():
                if name.startswith("word/media/"):
                    count += 1
                    if count > cap:
                        return count
        return count
    except Exception:
        logger.warning("Failed to count embedded images in docx", exc_info=True)
        return 0
    finally:
        if start_pos is not None:
            try:
                file.seek(start_pos)
            except Exception:
                pass


def read_docx_file(
    file: IO[Any],
    file_name: str = "",
    extract_images: bool = False,
    image_callback: Callable[[bytes, str], None] | None = None,
) -> tuple[str, Sequence[tuple[bytes, str]]]:
    """
    Extract text from a docx.
    Return (text_content, list_of_images).

    The caller can choose to provide a callback to handle images with the intent
    of avoiding materializing the list of images in memory.
    The images list returned is empty in this case.
    """
    md = get_markitdown_converter()
    from markitdown import (
        FileConversionException,
        StreamInfo,
        UnsupportedFormatException,
    )

    try:
        doc = md.convert(
            to_bytesio(file), stream_info=StreamInfo(mimetype=WORD_PROCESSING_MIME_TYPE)
        )
    except (
        BadZipFile,
        ValueError,
        FileConversionException,
        UnsupportedFormatException,
    ) as e:
        logger.warning(
            "Failed to extract docx %s: %s. Attempting to read as text file.",
            file_name or "docx file",
            e,
        )

        # May be an invalid docx, but still a valid text file
        file.seek(0)
        encoding = detect_encoding(file)
        text_content_raw, _ = read_text_file(
            file, encoding=encoding, ignore_onyx_metadata=False
        )
        return text_content_raw or "", []

    file.seek(0)

    if extract_images:
        if image_callback is None:
            return doc.markdown, list(extract_docx_images(to_bytesio(file)))
        # If a callback is provided, iterate and stream images without accumulating
        try:
            for img_file_bytes, img_file_name in extract_docx_images(to_bytesio(file)):
                image_callback(img_file_bytes, img_file_name)
        except Exception:
            logger.exception("Failed to stream docx images")
    return doc.markdown, []


def doc_to_text(file: IO[Any], file_name: str = "") -> str:
    """Extract text from a legacy binary Word (.doc) file.

    Uses the antiword system package (see backend/Dockerfile) in a subprocess —
    crashes and hangs in the C parser are contained by process isolation and
    the timeout instead of taking down the worker.
    """
    data = file.read()
    try:
        result = subprocess.run(
            ["antiword", "-m", "UTF-8.txt", "-"],
            input=data,
            capture_output=True,
            timeout=DOC_TEXT_EXTRACTION_TIMEOUT_SECONDS,
            check=False,
        )
    except FileNotFoundError as e:
        raise RuntimeError(
            "antiword is required to process .doc files but is not installed"
        ) from e
    except subprocess.TimeoutExpired as e:
        raise RuntimeError(
            f"antiword timed out after {DOC_TEXT_EXTRACTION_TIMEOUT_SECONDS}s "
            f"on {file_name or 'doc file'}"
        ) from e

    if result.returncode != 0:
        stderr = result.stderr.decode("utf-8", errors="replace").strip()
        raise RuntimeError(
            f"antiword failed on {file_name or 'doc file'} "
            f"(exit {result.returncode}): {stderr[:500]}"
        )
    return result.stdout.decode("utf-8", errors="replace")


# Raster signatures carved out of legacy .doc OLE streams; the vision captioning
# pipeline consumes PNG/JPEG, so vector EMF/WMF blips are not extracted.
_DOC_PNG_SIGNATURE = b"\x89PNG\r\n\x1a\n"
_DOC_JPEG_SIGNATURE = b"\xff\xd8\xff"
_DOC_PNG_END = b"IEND"
_DOC_JPEG_END = b"\xff\xd9"
_DOC_MIN_IMAGE_BYTES = 32


def _carve_raster_images(data: bytes) -> Iterator[tuple[bytes, str]]:
    """Yield PNG/JPEG blobs found anywhere in `data` by signature, deduped."""
    seen: set[bytes] = set()
    png_count = 0
    jpeg_count = 0
    pos = 0
    while pos < len(data):
        png_start = data.find(_DOC_PNG_SIGNATURE, pos)
        jpeg_start = data.find(_DOC_JPEG_SIGNATURE, pos)
        if png_start == -1 and jpeg_start == -1:
            return
        if png_start != -1 and (jpeg_start == -1 or png_start < jpeg_start):
            iend = data.find(_DOC_PNG_END, png_start)
            if iend == -1:
                pos = png_start + len(_DOC_PNG_SIGNATURE)
                continue
            end = min(iend + len(_DOC_PNG_END) + 4, len(data))
            blob = data[png_start:end]
            pos = end
            if len(blob) < _DOC_MIN_IMAGE_BYTES or blob in seen:
                continue
            seen.add(blob)
            png_count += 1
            yield blob, f"image_{png_count}.png"
        else:
            end = data.find(_DOC_JPEG_END, jpeg_start + len(_DOC_JPEG_SIGNATURE))
            if end == -1:
                pos = jpeg_start + len(_DOC_JPEG_SIGNATURE)
                continue
            end += len(_DOC_JPEG_END)
            blob = data[jpeg_start:end]
            pos = end
            if len(blob) < _DOC_MIN_IMAGE_BYTES or blob in seen:
                continue
            seen.add(blob)
            jpeg_count += 1
            yield blob, f"image_{jpeg_count}.jpg"


def extract_doc_images(doc_bytes: IO[Any]) -> Iterator[tuple[bytes, str]]:
    """Extract raster images embedded in a legacy binary Word (.doc) file.

    Word stores picture blips inside OLE streams (typically `Data`); unlike
    docx there is no manifest, so rasters are carved out of every stream by
    signature and deduped.
    """
    import olefile

    try:
        data = doc_bytes.read()
        ole = olefile.OleFileIO(BytesIO(data))
    except Exception:
        logger.exception("Failed to open .doc OLE structure for image extraction")
        return
    try:
        for stream_path in ole.listdir():
            try:
                stream_data = ole.openstream(stream_path).read()
            except Exception:
                continue
            for blob, name in _carve_raster_images(stream_data):
                yield blob, name
    except Exception:
        logger.exception("Failed to extract all doc images")
    finally:
        ole.close()


def count_doc_embedded_images(file: IO[Any], cap: int) -> int:
    """Return the number of raster images embedded in a .doc, short-circuiting
    at cap+1. Mirrors count_docx_embedded_images; always restores the file
    pointer."""

    try:
        start_pos = file.tell()
    except Exception:
        start_pos = None
    try:
        if start_pos is not None:
            file.seek(0)
        count = 0
        for _blob, _name in extract_doc_images(file):
            count += 1
            if count > cap:
                return count
        return count
    except Exception:
        logger.warning("Failed to count embedded images in doc", exc_info=True)
        return 0
    finally:
        if start_pos is not None:
            file.seek(start_pos)


def read_doc_file(
    file: IO[Any],
    file_name: str = "",
    extract_images: bool = False,
    image_callback: Callable[[bytes, str], None] | None = None,
) -> tuple[str, Sequence[tuple[bytes, str]]]:
    """
    Extract text and optionally images from a legacy binary Word (.doc) file.
    Return (text_content, list_of_images).
    """
    text = doc_to_text(file, file_name=file_name)

    file.seek(0)

    if extract_images:
        if image_callback is None:
            return text, list(extract_doc_images(file))
        try:
            for img_file_bytes, img_file_name in extract_doc_images(file):
                image_callback(img_file_bytes, img_file_name)
        except Exception:
            logger.exception("Failed to stream doc images")
    return text, []


def _is_odf_raster_entry(name: str) -> bool:
    """ODF pictures are referenced from the manifest and may sit at any path
    (LibreOffice writes `Pictures/…`); the document thumbnail preview must not
    be picked up."""
    lowered = name.lower()
    return lowered.startswith("pictures/") or (
        lowered.endswith((".png", ".jpg", ".jpeg"))
        and not lowered.startswith(("thumbnails/", "meta/", "styles/"))
    )


def extract_odt_images(odt_bytes: IO[Any]) -> Iterator[tuple[bytes, str]]:
    """Extract raster images embedded in an OpenDocument Text (.odt) file.

    Only formats the vision captioning pipeline consumes (PNG/JPEG) are
    yielded.
    """
    try:
        with zipfile.ZipFile(odt_bytes) as z:
            for name in z.namelist():
                if _is_odf_raster_entry(name):
                    yield (z.read(name), name.split("/")[-1])
    except Exception:
        logger.exception("Failed to extract all odt images")


def count_odt_embedded_images(file: IO[Any], cap: int) -> int:
    """Return the number of raster images embedded in an .odt, short-circuiting
    at cap+1. Mirrors count_docx_embedded_images; always restores the file
    pointer."""
    try:
        start_pos = file.tell()
    except Exception:
        start_pos = None
    try:
        if start_pos is not None:
            file.seek(0)
        count = 0
        with zipfile.ZipFile(file) as z:
            for name in z.namelist():
                if _is_odf_raster_entry(name):
                    count += 1
                    if count > cap:
                        return count
        return count
    except Exception:
        logger.warning("Failed to count embedded images in odt", exc_info=True)
        return 0
    finally:
        if start_pos is not None:
            file.seek(start_pos)


def read_odt_file(
    file: IO[Any],
    file_name: str = "",
    extract_images: bool = False,
    image_callback: Callable[[bytes, str], None] | None = None,
) -> tuple[str, Sequence[tuple[bytes, str]]]:
    """
    Extract text and optionally images from an OpenDocument Text (.odt) file.
    Paragraphs and headings become lines; tables are rendered as
    pipe-separated rows so column structure survives indexing.
    Return (text_content, list_of_images).
    """
    from odf import opendocument, teletype
    from odf import table as odf_table
    from odf.element import Element
    from odf.namespaces import TABLENS, TEXTNS

    PARAGRAPH_QNAMES = {(TEXTNS, "p"), (TEXTNS, "h")}
    TABLE_QNAME = (TABLENS, "table")
    CELL_QNAMES = {(TABLENS, "table-cell"), (TABLENS, "covered-table-cell")}

    try:
        document = opendocument.load(file)
    except Exception as e:
        raise RuntimeError(
            f"Failed to parse ODT file {file_name or 'odt file'}: {e}"
        ) from e

    paragraphs: list[str] = []

    def render_table(table: Element) -> str:
        rows: list[str] = []
        for row in table.getElementsByType(odf_table.TableRow):
            cells: list[str] = []
            for cell in row.childNodes:
                if cell.qname not in CELL_QNAMES:
                    continue
                cell_text = teletype.extractText(cell).strip()
                repeat = int(cell.getAttribute("numbercolumnsrepeated") or 1)
                # Content cells repeat for column alignment; empty repeats are
                # usually trailing padding. Cap so a 1024-column repeat can't
                # blow up the output.
                cells.extend([cell_text] * min(repeat, MAX_ODS_CELL_REPEAT))
            rows.append(" | ".join(cells))
        return "\n".join(rows)

    def walk(node: Element) -> None:
        for child in node.childNodes:
            if not isinstance(child, Element):
                continue
            if child.qname in PARAGRAPH_QNAMES:
                text = teletype.extractText(child).strip()
                if text:
                    paragraphs.append(text)
            elif child.qname == TABLE_QNAME:
                rendered = render_table(child)
                if rendered:
                    paragraphs.append(rendered)
            else:
                walk(child)

    walk(document.text)
    text_content = "\n".join(paragraphs)

    file.seek(0)

    if extract_images:
        if image_callback is None:
            return text_content, list(extract_odt_images(file))
        try:
            for img_file_bytes, img_file_name in extract_odt_images(file):
                image_callback(img_file_bytes, img_file_name)
        except Exception:
            logger.exception("Failed to stream odt images")
    return text_content, []


def extract_pptx_images(pptx_bytes: IO[Any]) -> Iterator[tuple[bytes, str]]:
    """
    Given the bytes of a pptx file, extract all the images.
    Returns an iterator of tuples (image_bytes, image_name).
    """
    try:
        with zipfile.ZipFile(pptx_bytes) as z:
            for name in z.namelist():
                if name.startswith("ppt/media/"):
                    yield (z.read(name), name.split("/")[-1])
    except Exception:
        logger.exception("Failed to extract all pptx images")


def pptx_to_text(file: IO[Any], file_name: str = "") -> str:
    md = get_markitdown_converter()
    from markitdown import (
        FileConversionException,
        StreamInfo,
        UnsupportedFormatException,
    )

    stream_info = StreamInfo(
        mimetype=PRESENTATION_MIME_TYPE, filename=file_name or None, extension=".pptx"
    )
    try:
        presentation = md.convert(to_bytesio(file), stream_info=stream_info)
    except (
        BadZipFile,
        ValueError,
        FileConversionException,
        UnsupportedFormatException,
    ) as e:
        error_str = f"Failed to extract text from {file_name or 'pptx file'}: {e}"
        logger.warning(error_str)
        return ""
    return presentation.markdown


def read_pptx_file(
    file: IO[Any],
    file_name: str = "",
    extract_images: bool = False,
    image_callback: Callable[[bytes, str], None] | None = None,
) -> tuple[str, Sequence[tuple[bytes, str]]]:
    """
    Extract text and optionally images from a pptx.
    Return (text_content, list_of_images).
    """
    text = pptx_to_text(file, file_name=file_name)

    file.seek(0)

    if extract_images:
        if image_callback is None:
            return text, list(extract_pptx_images(to_bytesio(file)))
        try:
            for img_file_bytes, img_file_name in extract_pptx_images(to_bytesio(file)):
                image_callback(img_file_bytes, img_file_name)
        except Exception:
            logger.exception("Failed to stream pptx images")
    return text, []


def _columns_to_keep(col_has_data: bytearray, max_empty: int) -> list[int]:
    """Keep non-empty columns, plus runs of up to ``max_empty`` empty columns
    between them. Trailing empty columns are dropped."""
    kept: list[int] = []
    empty_buffer: list[int] = []
    for c, has in enumerate(col_has_data):
        if has:
            kept.extend(empty_buffer[:max_empty])
            kept.append(c)
            empty_buffer = []
        else:
            empty_buffer.append(c)
    return kept


def _sheet_to_csv(rows: Iterator[tuple[Any, ...]]) -> str:
    """Stream worksheet rows into CSV text without materializing a dense matrix.

    Empty rows are never stored. Column occupancy is tracked as a ``bytearray``
    bitmap so column trimming needs no transpose or copy. Runs of empty
    rows/columns longer than 2 are collapsed; shorter runs are preserved.

    Scanning stops once ``MAX_XLSX_CELLS_PER_SHEET`` non-empty cells have been
    seen; the output gets a truncation marker row appended so downstream
    indexing sees that the sheet was cut off.
    """
    MAX_EMPTY_ROWS_IN_OUTPUT = 2
    MAX_EMPTY_COLS_IN_OUTPUT = 2
    TRUNCATION_MARKER = "[truncated: sheet exceeded cell limit]"

    non_empty_rows: list[tuple[int, list[str]]] = []
    col_has_data = bytearray()
    total_non_empty = 0
    truncated = False

    for row_idx, row_vals in enumerate(rows):
        # Fast-reject empty rows before allocating a list of "".
        if not any(v is not None and v != "" for v in row_vals):
            continue

        cells = ["" if v is None else str(v) for v in row_vals]
        non_empty_rows.append((row_idx, cells))

        if len(cells) > len(col_has_data):
            col_has_data.extend(b"\x00" * (len(cells) - len(col_has_data)))
        for i, v in enumerate(cells):
            if v:
                col_has_data[i] = 1
                total_non_empty += 1

        if total_non_empty > MAX_XLSX_CELLS_PER_SHEET:
            truncated = True
            break

    if not non_empty_rows:
        return ""

    keep_cols = _columns_to_keep(col_has_data, MAX_EMPTY_COLS_IN_OUTPUT)
    if not keep_cols:
        return ""

    buf = io.StringIO()
    writer = csv.writer(buf, lineterminator="\n")
    blank_row = [""] * len(keep_cols)
    last_idx = -1
    for row_idx, cells in non_empty_rows:
        gap = row_idx - last_idx - 1
        if gap > 0:
            for _ in range(min(gap, MAX_EMPTY_ROWS_IN_OUTPUT)):
                writer.writerow(blank_row)
        writer.writerow([cells[c] if c < len(cells) else "" for c in keep_cols])
        last_idx = row_idx

    if truncated:
        writer.writerow([TRUNCATION_MARKER])

    return buf.getvalue().rstrip("\n")


def _load_readonly_workbook(file: IO[Any], file_name: str) -> openpyxl.Workbook | None:
    """Load a read-only workbook, returning None (and logging) for the BadZipFile
    / known-openpyxl-bug cases the xlsx indexers treat as skip-and-continue rather
    than a hard failure."""
    try:
        return openpyxl.load_workbook(file, read_only=True, data_only=True)
    except BadZipFile as e:
        error_str = f"Failed to extract text from {file_name or 'xlsx file'}: {e}"
        if file_name.startswith("~"):
            logger.debug(error_str + " (this is expected for files with ~)")
        else:
            logger.warning(error_str)
        return None
    except Exception as e:
        if any(s in str(e) for s in KNOWN_OPENPYXL_BUGS):
            logger.warning(
                "Failed to extract text from %s. This happens due to a bug in openpyxl. %s",
                file_name or "xlsx file",
                e,
            )
            return None
        raise


def xlsx_sheet_extraction(file: IO[Any], file_name: str = "") -> list[tuple[str, str]]:
    """
    Converts each sheet in the excel file to a csv condensed string.
    Returns a string and the worksheet title for each worksheet

    Returns a list of (csv_text, sheet)
    """
    workbook = _load_readonly_workbook(file, file_name)
    if workbook is None:
        return []

    sheets: list[tuple[str, str]] = []
    try:
        for sheet in workbook.worksheets:
            # Declared dimensions can be different to what is actually there
            ro_sheet = cast(ReadOnlyWorksheet, sheet)
            ro_sheet.reset_dimensions()
            csv_text = _sheet_to_csv(ro_sheet.iter_rows(values_only=True))
            sheets.append((csv_text.strip(), ro_sheet.title))
    finally:
        workbook.close()

    return sheets


class StreamedSheet(NamedTuple):
    """One worksheet rendered to CSV, staged in the file store, and referenced
    by `csv_file_id`."""

    title: str
    csv_file_id: str


def _row_has_content(row: tuple[Any, ...]) -> bool:
    return any(v is not None and v != "" for v in row)


def _cell(value: Any) -> str:
    return "" if value is None else str(value)


def stage_xlsx_sheets(
    file: IO[bytes],
    stage: Callable[[IO[bytes], str], str],
    file_name: str = "",
) -> list[StreamedSheet]:
    """Stream each non-empty worksheet to a temp CSV row by row (never holding a
    full sheet in memory), then stage it via `stage` and reference it by
    `csv_file_id`. Empty rows are dropped; columns are not trimmed."""
    sheets: list[StreamedSheet] = []
    workbook = _load_readonly_workbook(file, file_name)
    if workbook is None:
        return sheets
    try:
        for sheet in workbook.worksheets:
            ro_sheet = cast(ReadOnlyWorksheet, sheet)
            ro_sheet.reset_dimensions()
            with tempfile.TemporaryFile(mode="w+", encoding="utf-8", newline="") as tmp:
                writer = csv.writer(tmp, lineterminator="\n")
                for row in ro_sheet.iter_rows(values_only=True):
                    if _row_has_content(row):
                        writer.writerow([_cell(v) for v in row])
                tmp.flush()
                binary = cast(IO[bytes], tmp.buffer)
                if binary.seek(0, io.SEEK_END) == 0:
                    continue
                binary.seek(0)
                sheets.append(StreamedSheet(ro_sheet.title, stage(binary, "text/csv")))
    finally:
        workbook.close()
    return sheets


def xlsx_to_text(file: IO[Any], file_name: str = "") -> str:
    sheets = xlsx_sheet_extraction(file, file_name)
    return TEXT_SECTION_SEPARATOR.join(
        csv_text for csv_text, _title in sheets if csv_text
    )


def _xls_rows(sheet: Any, book: Any) -> Iterator[tuple[Any, ...]]:
    """Yield one tuple of cell values per row, converting Excel date serials
    (xlrd returns floats for them) to datetime strings, and integral floats to
    ints, so the CSV matches what the openpyxl path produces."""
    import xlrd

    for r in range(sheet.nrows):
        values: list[Any] = []
        for cell in sheet.row(r):
            if cell.ctype == xlrd.XL_CELL_DATE:
                values.append(
                    xlrd.xldate_as_datetime(cell.value, book.datemode).isoformat()
                )
            elif cell.ctype == xlrd.XL_CELL_NUMBER and float(cell.value).is_integer():
                values.append(int(cell.value))
            else:
                values.append(cell.value)
        yield tuple(values)


def xls_sheet_extraction(file: IO[Any], file_name: str = "") -> list[tuple[str, str]]:
    """Converts each sheet in a legacy binary Excel (.xls) file to a condensed
    CSV string. Returns (csv_text, sheet_title) per worksheet."""
    import xlrd

    file_contents = file.read()
    try:
        book = xlrd.open_workbook(file_contents=file_contents)
    except Exception as e:
        error_str = f"Failed to extract text from {file_name or 'xls file'}: {e}"
        logger.warning(error_str)
        return []

    sheets: list[tuple[str, str]] = []
    for sheet in book.sheets():
        csv_text = _sheet_to_csv(_xls_rows(sheet, book))
        sheets.append((csv_text.strip(), sheet.name))
    return sheets


def xls_to_text(file: IO[Any], file_name: str = "") -> str:
    sheets = xls_sheet_extraction(file, file_name)
    return TEXT_SECTION_SEPARATOR.join(
        csv_text for csv_text, _title in sheets if csv_text
    )


def _ods_cell_value(cell: Any) -> Any:
    """Value of an ODF table cell, mirroring what the openpyxl path yields:
    numeric/date cells come from their office:value attributes, text cells
    from their text content."""
    from odf import teletype

    value_type = cell.getAttribute("valuetype")
    if value_type in ("float", "percentage", "currency"):
        raw_value = cell.getAttribute("value")
        if raw_value is not None:
            number = float(raw_value)
            return int(number) if number.is_integer() else number
    if value_type == "date":
        date_value = cell.getAttribute("datevalue")
        if date_value is not None:
            return date_value
    return teletype.extractText(cell)


def ods_sheet_extraction(file: IO[Any], file_name: str = "") -> list[tuple[str, str]]:
    """Converts each table in an OpenDocument Spreadsheet (.ods) file to a
    condensed CSV string. Returns (csv_text, table_name) per table."""
    from odf import opendocument
    from odf import table as odf_table
    from odf.element import Element
    from odf.namespaces import TABLENS

    CELL_QNAMES = {(TABLENS, "table-cell"), (TABLENS, "covered-table-cell")}

    try:
        document = opendocument.load(file)
    except Exception as e:
        error_str = f"Failed to extract text from {file_name or 'ods file'}: {e}"
        logger.warning(error_str)
        return []

    def table_rows(table: Element) -> Iterator[tuple[Any, ...]]:
        for row in table.getElementsByType(odf_table.TableRow):
            values: list[Any] = []
            for cell in row.childNodes:
                if cell.qname not in CELL_QNAMES:
                    continue
                values.extend(
                    [_ods_cell_value(cell)]
                    * min(
                        int(cell.getAttribute("numbercolumnsrepeated") or 1),
                        MAX_ODS_CELL_REPEAT,
                    )
                )
            yield tuple(values)

    sheets: list[tuple[str, str]] = []
    for table in document.spreadsheet.getElementsByType(odf_table.Table):
        title = table.getAttribute("name") or ""
        csv_text = _sheet_to_csv(table_rows(table))
        sheets.append((csv_text.strip(), title))
    return sheets


def ods_to_text(file: IO[Any], file_name: str = "") -> str:
    sheets = ods_sheet_extraction(file, file_name)
    return TEXT_SECTION_SEPARATOR.join(
        csv_text for csv_text, _title in sheets if csv_text
    )


def spreadsheet_sheet_extraction(
    file: IO[Any], file_name: str = ""
) -> list[tuple[str, str]]:
    """Extension-dispatching sheet extraction covering every supported
    spreadsheet format (.xlsx/.xlsm, .xls, .ods). Returns (csv_text,
    sheet_title) per worksheet."""
    extension = get_file_ext(file_name)
    if extension == ".xls":
        return xls_sheet_extraction(file, file_name)
    if extension == ".ods":
        return ods_sheet_extraction(file, file_name)
    return xlsx_sheet_extraction(file, file_name)


def stage_spreadsheet_sheets(
    file: IO[bytes],
    stage: Callable[[IO[bytes], str], str],
    file_name: str = "",
) -> list[StreamedSheet]:
    """Extension-dispatching counterpart of stage_xlsx_sheets: stream each
    worksheet to a staged CSV for the code sandbox. .xlsx/.xlsm use the
    openpyxl row streaming; .xls/.ods extract first and stage per sheet."""
    extension = get_file_ext(file_name)
    if extension not in (".xls", ".ods"):
        return stage_xlsx_sheets(file, stage, file_name=file_name)

    sheets: list[StreamedSheet] = []
    for csv_text, title in spreadsheet_sheet_extraction(file, file_name):
        if not csv_text:
            continue
        sheets.append(
            StreamedSheet(title, stage(BytesIO(csv_text.encode("utf-8")), "text/csv"))
        )
    return sheets


def eml_to_text(file: IO[Any]) -> str:
    encoding = detect_encoding(file)
    text_file = io.TextIOWrapper(file, encoding=encoding)
    parser = EmailParser()
    try:
        message = parser.parse(text_file)
    finally:
        try:
            # Keep underlying upload handle open for downstream consumers.
            raw_file = text_file.detach()
        except Exception as detach_error:
            logger.warning(
                "Failed to detach TextIOWrapper for EML upload, using original file: %s",
                detach_error,
            )
            raw_file = file
        try:
            raw_file.seek(0)
        except Exception:
            pass

    text_content = []
    for part in message.walk():
        if part.get_content_type().startswith("text/plain"):
            payload = part.get_payload()
            if isinstance(payload, str):
                text_content.append(payload)
            elif isinstance(payload, list):
                text_content.extend(item for item in payload if isinstance(item, str))
            else:
                logger.warning("Unexpected payload type: %s", type(payload))
    return TEXT_SECTION_SEPARATOR.join(text_content)


def epub_to_text(file: IO[Any]) -> str:
    with zipfile.ZipFile(file) as epub:
        text_content = []
        for item in epub.infolist():
            if item.filename.endswith(".xhtml") or item.filename.endswith(".html"):
                with epub.open(item) as html_file:
                    text_content.append(parse_html_page_basic(html_file))
        return TEXT_SECTION_SEPARATOR.join(text_content)


def file_io_to_text(file: IO[Any]) -> str:
    encoding = detect_encoding(file)
    file_content, _ = read_text_file(file, encoding=encoding)
    return file_content


def extract_file_text(
    file: IO[Any],
    file_name: str,
    break_on_unprocessable: bool = True,
    extension: str | None = None,
) -> str:
    """
    Legacy function that returns *only text*, ignoring embedded images.
    For backward-compatibility in code that only wants text.

    NOTE: Ignoring seems to be defined as returning an empty string for files it can't
    handle (such as images).
    """
    try:
        if get_unstructured_api_key():
            try:
                return unstructured_to_text(file, file_name)
            except Exception as unstructured_error:
                logger.error(
                    "Failed to process with Unstructured: %s. Falling back to normal processing.",
                    str(unstructured_error),
                )
        return extract_file_text_locally(file, file_name, extension)

    except Exception as e:
        if break_on_unprocessable:
            raise RuntimeError(
                f"Failed to process file {file_name or 'Unknown'}: {str(e)}"
            ) from e
        logger.warning("Failed to process file %s: %s", file_name or "Unknown", str(e))
        return ""


def extract_file_text_locally(
    file: IO[Any],
    file_name: str,
    extension: str | None = None,
    isolate_pdfium: bool = True,
) -> str:
    """Text by extension from the in-process parsers only.

    Never reaches the database or Redis (the Unstructured key lives there), so
    it can run in a child process that has neither. Raises on a file no parser
    accepts. See ``read_pdf_file`` for ``isolate_pdfium``.
    """
    extension_to_function: dict[str, Callable[[IO[Any]], str]] = {
        ".pdf": lambda f: pdf_to_text(f, isolate_pdfium=isolate_pdfium),
        ".doc": lambda f: read_doc_file(f)[0],  # no images
        ".docx": lambda f: read_docx_file(f, file_name)[0],  # no images
        ".odt": lambda f: read_odt_file(f, file_name)[0],  # no images
        ".pptx": lambda f: pptx_to_text(f, file_name),
        ".xlsx": lambda f: xlsx_to_text(f, file_name),
        # openpyxl reads macro-enabled workbooks like any other.
        ".xlsm": lambda f: xlsx_to_text(f, file_name),
        ".xls": lambda f: xls_to_text(f, file_name),
        ".ods": lambda f: ods_to_text(f, file_name),
        ".eml": eml_to_text,
        ".epub": epub_to_text,
        ".html": parse_html_page_basic,
    }
    if extension is None:
        extension = get_file_ext(file_name)

    if extension in OnyxFileExtensions.TEXT_AND_DOCUMENT_EXTENSIONS:
        func = extension_to_function.get(extension, file_io_to_text)
        file.seek(0)
        return func(file)

    # If unknown extension, maybe it's a text file
    file.seek(0)
    if is_text_file(file):
        return file_io_to_text(file)

    raise ValueError("Unknown file extension or not recognized as text data")


class ExtractionResult(NamedTuple):
    """Structured result from text and image extraction from various file types."""

    text_content: str
    embedded_images: Sequence[tuple[bytes, str]]
    metadata: dict[str, Any]


def extract_result_from_text_file(file: IO[Any]) -> ExtractionResult:
    encoding = detect_encoding(file)
    text_content_raw, file_metadata = read_text_file(
        file, encoding=encoding, ignore_onyx_metadata=False
    )
    return ExtractionResult(
        text_content=text_content_raw,
        embedded_images=[],
        metadata=file_metadata,
    )


def extract_text_and_images(
    file: IO[Any],
    file_name: str,
    pdf_pass: str | None = None,
    content_type: str | None = None,
    image_callback: Callable[[bytes, str], None] | None = None,
) -> ExtractionResult:
    """
    Primary new function for the updated connector.
    Returns structured extraction result with text content, embedded images, and metadata.

    Args:
        file: File-like object to extract content from.
        file_name: Name of the file (used to determine extension/type).
        pdf_pass: Optional password for encrypted PDFs.
        content_type: Optional MIME type override for the file.
        image_callback: Optional callback for streaming image extraction. When provided,
            embedded images are passed to this callback one at a time as (bytes, filename)
            instead of being accumulated in the returned ExtractionResult.embedded_images
            list. This is a memory optimization for large documents with many images -
            the caller can process/store each image immediately rather than holding all
            images in memory. When using a callback, ExtractionResult.embedded_images
            will be an empty list.

    Returns:
        ExtractionResult containing text_content, embedded_images (empty if callback used),
        and metadata extracted from the file.
    """
    res = _extract_text_and_images(
        file, file_name, pdf_pass, content_type, image_callback
    )
    # Clean up any temporary objects and force garbage collection
    unreachable = gc.collect()
    logger.info("Unreachable objects: %s", unreachable)

    return res


def _extract_text_and_images(
    file: IO[Any],
    file_name: str,
    pdf_pass: str | None = None,
    content_type: str | None = None,
    image_callback: Callable[[bytes, str], None] | None = None,
) -> ExtractionResult:
    file.seek(0)

    if get_unstructured_api_key():
        try:
            text_content = unstructured_to_text(file, file_name)
            return ExtractionResult(
                text_content=text_content, embedded_images=[], metadata={}
            )
        except Exception as e:
            logger.error(
                "Failed to process with Unstructured: %s. Falling back to normal processing.",
                str(e),
            )
            file.seek(0)  # Reset file pointer just in case

    # When we upload a document via a connector or MyDocuments, we extract and store the content of files
    # with content types in UploadMimeTypes.DOCUMENT_MIME_TYPES as plain text files.
    # As a result, the file name extension may differ from the original content type.
    # We process files with a plain text content type first to handle this scenario.
    if content_type in OnyxMimeTypes.TEXT_MIME_TYPES:
        return extract_result_from_text_file(file)

    # Default processing
    try:
        extension = get_file_ext(file_name)
        # One setting read per file. With extraction off, no image is decoded
        # for any format, which is what lets upload validation skip its caps.
        extract_images = get_image_extraction_and_analysis_enabled()

        if extension == ".docx":
            text_content, images = read_docx_file(
                file,
                file_name,
                extract_images=extract_images,
                image_callback=image_callback,
            )
            return ExtractionResult(
                text_content=text_content, embedded_images=images, metadata={}
            )

        if extension == ".pdf":
            text_content, pdf_metadata, images = read_pdf_file(
                file,
                pdf_pass,
                extract_images=extract_images,
                image_callback=image_callback,
            )
            return ExtractionResult(
                text_content=text_content, embedded_images=images, metadata=pdf_metadata
            )

        if extension == ".pptx":
            text_content, images = read_pptx_file(
                file,
                file_name,
                extract_images=extract_images,
                image_callback=image_callback,
            )
            return ExtractionResult(
                text_content=text_content, embedded_images=images, metadata={}
            )

        if extension == ".doc":
            text_content, images = read_doc_file(
                file,
                file_name,
                extract_images=extract_images,
                image_callback=image_callback,
            )
            return ExtractionResult(
                text_content=text_content, embedded_images=images, metadata={}
            )

        if extension == ".odt":
            text_content, images = read_odt_file(
                file,
                file_name,
                extract_images=extract_images,
                image_callback=image_callback,
            )
            return ExtractionResult(
                text_content=text_content, embedded_images=images, metadata={}
            )

        if extension in (".xlsx", ".xlsm", ".xls", ".ods"):
            sheets = spreadsheet_sheet_extraction(file, file_name)
            return ExtractionResult(
                text_content=TEXT_SECTION_SEPARATOR.join(
                    csv_text for csv_text, _title in sheets if csv_text
                ),
                embedded_images=[],
                metadata={},
            )

        if extension == ".eml":
            return ExtractionResult(
                text_content=eml_to_text(file), embedded_images=[], metadata={}
            )

        if extension == ".epub":
            return ExtractionResult(
                text_content=epub_to_text(file), embedded_images=[], metadata={}
            )

        if extension == ".html":
            return ExtractionResult(
                text_content=parse_html_page_basic(file),
                embedded_images=[],
                metadata={},
            )

        # If we reach here and it's a recognized text extension
        if extension in OnyxFileExtensions.PLAIN_TEXT_EXTENSIONS:
            return extract_result_from_text_file(file)

        # If it's an image file or something else, we do not parse embedded images from them
        # just return empty text
        return ExtractionResult(text_content="", embedded_images=[], metadata={})

    except Exception as e:
        logger.exception("Failed to extract text/images from %s: %s", file_name, e)
        return ExtractionResult(text_content="", embedded_images=[], metadata={})


def docx_to_txt_filename(file_path: str) -> str:
    return file_path.rsplit(".", 1)[0] + ".txt"
